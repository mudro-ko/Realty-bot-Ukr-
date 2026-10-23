"""Excel-звіт з архіву: аркуш на кожен пошук (з посиланнями) + аркуш «Аналітика»."""
import re
import statistics
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(bold=True, size=11, color="1F4E78")
LINK_FONT = Font(color="0563C1", underline="single")
DROP_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="D9D9D9")
COND_SHORT = {"needs_repair": "під ремонт", "cosmetic": "житловий/косметика", "renovated": "з ремонтом", "unknown": "—"}


def _sheet_title(name, used):
    t = re.sub(r"[\[\]\*\?/\\:]", " ", name)[:28].strip() or "Пошук"
    base, n = t, 2
    while t in used:
        t = f"{base[:25]} {n}"
        n += 1
    used.add(t)
    return t


def _ppm(r):
    if r.get("price_usd") and r.get("area"):
        return round(r["price_usd"] / r["area"])
    return None


def _med(vals):
    vals = [v for v in vals if v]
    return round(statistics.median(vals)) if vals else None


def _header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(ws.column_dimensions[get_column_letter(i)].width or 0, w)


LIST_COLS = ["Знайдено", "Ціна $", "Ціна ₴", "$/м²", "Кімн.", "Площа м²", "Поверх", "Адреса",
             "Стан", "Власник", "Зміна ціни", "Сайт", "Посилання", "Опис"]
LIST_W = [16, 11, 12, 8, 7, 10, 9, 42, 18, 9, 14, 13, 12, 60]


def listing_sheet(ws, name, rows, deal):
    ws.append([name])
    ws["A1"].font = TITLE_FONT
    ws.append([f"Оголошень в архіві: {len(rows)} · сформовано {datetime.now():%d.%m.%Y %H:%M}"])
    _header(ws, 3, LIST_COLS, LIST_W)
    for r in rows:
        if deal == "sale":
            delta = (r["price_usd"] - r["first_price_usd"]) if r.get("price_usd") and r.get("first_price_usd") else 0
            delta_txt = f"{delta:+,}$".replace(",", " ") if delta else ""
        else:
            delta = (r["price_uah"] - r["first_price_uah"]) if r.get("price_uah") and r.get("first_price_uah") else 0
            delta_txt = f"{delta:+,}₴".replace(",", " ") if delta else ""
        floor = f"{r['floor']}/{r['floors']}" if r.get("floor") and r.get("floors") else (r.get("floor") or "")
        ws.append([
            r.get("first_seen"), r.get("price_usd"), r.get("price_uah"), _ppm(r) if deal == "sale" else None,
            r.get("rooms"), r.get("area"), floor, r.get("geo") or r.get("street"),
            COND_SHORT.get(r.get("condition"), "—"), "так" if r.get("is_owner") else "",
            delta_txt, r.get("origin"), "відкрити", " ".join((r.get("text") or "").split())[:300],
        ])
        i = ws.max_row
        link = ws.cell(row=i, column=13)
        link.hyperlink, link.font = r.get("url"), LINK_FONT
        if delta < 0:
            for c in range(1, len(LIST_COLS) + 1):
                ws.cell(row=i, column=c).fill = DROP_FILL
    for col in (2, 3, 4):
        for c in ws.iter_cols(min_col=col, max_col=col, min_row=4):
            for cell in c:
                cell.number_format = "#,##0"
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=14, max_col=14):
        for cell in row:
            cell.alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "B4"
    if rows:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(LIST_COLS))}{ws.max_row}"


def analytics_block(ws, start_row, name, rows, deal, link_sheet):
    """Повертає наступний вільний рядок."""
    r0 = start_row
    ws.cell(row=r0, column=1, value=name).font = TITLE_FONT
    ws.cell(row=r0, column=6, value="→ усі оголошення").hyperlink = f"#'{link_sheet}'!A1"
    ws.cell(row=r0, column=6).font = LINK_FONT
    money = "price_usd" if deal == "sale" else "price_uah"
    sym = "$" if deal == "sale" else "₴"
    prices = [x[money] for x in rows if x.get(money)]
    ppms = [_ppm(x) for x in rows] if deal == "sale" else []
    r = r0 + 1
    summary = [
        ("Оголошень", len(rows)),
        (f"Медіанна ціна, {sym}", _med(prices)),
        (f"Мінімальна ціна, {sym}", min(prices) if prices else None),
        (f"Максимальна ціна, {sym}", max(prices) if prices else None),
    ]
    if deal == "sale":
        summary.append(("Медіана $/м²", _med(ppms)))
    drops = [x for x in rows if x.get(money) and x.get("first_" + money) and x[money] < x["first_" + money]]
    summary.append(("Подешевшали з моменту появи", len(drops)))
    for k, v in summary:
        ws.cell(row=r, column=1, value=k)
        c = ws.cell(row=r, column=2, value=v)
        c.number_format = "#,##0"
        c.font = Font(bold=True)
        r += 1

    # за кімнатами
    r += 1
    ws.cell(row=r, column=1, value="За кількістю кімнат").font = SUB_FONT
    r += 1
    _header(ws, r, ["Кімнат", "Оголошень", f"Медіана, {sym}", "Медіана $/м²" if deal == "sale" else "Медіана м²"])
    by_rooms = defaultdict(list)
    for x in rows:
        by_rooms[x.get("rooms") or 0].append(x)
    rooms_first = r + 1
    for k in sorted(by_rooms):
        g = by_rooms[k]
        r += 1
        ws.cell(row=r, column=1, value=f"{k}к" if k else "?")
        ws.cell(row=r, column=2, value=len(g))
        ws.cell(row=r, column=3, value=_med([x.get(money) for x in g])).number_format = "#,##0"
        ws.cell(row=r, column=4, value=_med([_ppm(x) for x in g]) if deal == "sale" else _med([x.get("area") for x in g])).number_format = "#,##0"
    rooms_last = r

    # за станом
    r += 2
    ws.cell(row=r, column=1, value="За станом ремонту").font = SUB_FONT
    r += 1
    _header(ws, r, ["Стан", "Оголошень", f"Медіана, {sym}", "Медіана $/м²" if deal == "sale" else ""])
    by_cond = defaultdict(list)
    for x in rows:
        by_cond[x.get("condition") or "unknown"].append(x)
    for k in ["needs_repair", "cosmetic", "unknown", "renovated"]:
        if k not in by_cond:
            continue
        g = by_cond[k]
        r += 1
        ws.cell(row=r, column=1, value=COND_SHORT[k])
        ws.cell(row=r, column=2, value=len(g))
        ws.cell(row=r, column=3, value=_med([x.get(money) for x in g])).number_format = "#,##0"
        if deal == "sale":
            ws.cell(row=r, column=4, value=_med([_ppm(x) for x in g])).number_format = "#,##0"

    # динаміка по тижнях
    r += 2
    ws.cell(row=r, column=1, value="Динаміка по тижнях (за датою появи в моніторингу)").font = SUB_FONT
    r += 1
    _header(ws, r, ["Тиждень", "Нових", f"Медіана, {sym}", "Медіана $/м²" if deal == "sale" else ""])
    weeks = defaultdict(list)
    for x in rows:
        try:
            d = datetime.strptime((x.get("first_seen") or "")[:10], "%Y-%m-%d")
        except ValueError:
            continue
        y, w, _ = d.isocalendar()
        weeks[f"{y}-W{w:02d}"].append(x)
    wk_first = r + 1
    for k in sorted(weeks):
        g = weeks[k]
        r += 1
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=len(g))
        ws.cell(row=r, column=3, value=_med([x.get(money) for x in g])).number_format = "#,##0"
        if deal == "sale":
            ws.cell(row=r, column=4, value=_med([_ppm(x) for x in g])).number_format = "#,##0"
    wk_last = r

    # графіки праворуч від таблиць
    if rooms_last > rooms_first - 1 and by_rooms:
        ch = BarChart()
        ch.title, ch.height, ch.width = "Оголошень за кімнатами", 6, 12
        ch.add_data(Reference(ws, min_col=2, min_row=rooms_first - 1, max_row=rooms_last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=rooms_first, max_row=rooms_last))
        ch.legend = None
        ws.add_chart(ch, f"H{r0 + 1}")
    if wk_last >= wk_first:
        lc = LineChart()
        lc.title, lc.height, lc.width = ("Медіана $/м² по тижнях" if deal == "sale" else f"Медіана ціни, {sym}"), 6, 12
        col = 4 if deal == "sale" else 3
        lc.add_data(Reference(ws, min_col=col, min_row=wk_first - 1, max_row=wk_last), titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=1, min_row=wk_first, max_row=wk_last))
        lc.legend = None
        for ser in lc.series:
            ser.marker.symbol = "circle"
            ser.marker.size = 7
        ws.add_chart(lc, f"H{r0 + 14}")

    # топ найвигідніших
    r += 2
    title = "Топ-10 найдешевших за м²" if deal == "sale" else "Топ-10 найдешевших"
    ws.cell(row=r, column=1, value=title).font = SUB_FONT
    r += 1
    _header(ws, r, [f"Ціна, {sym}", "$/м²" if deal == "sale" else "Площа", "Кімн.", "Адреса", "Посилання"])
    key = (lambda x: _ppm(x) or 10 ** 9) if deal == "sale" else (lambda x: x.get(money) or 10 ** 12)
    for x in sorted(rows, key=key)[:10]:
        r += 1
        ws.cell(row=r, column=1, value=x.get(money)).number_format = "#,##0"
        ws.cell(row=r, column=2, value=_ppm(x) if deal == "sale" else x.get("area")).number_format = "#,##0"
        ws.cell(row=r, column=3, value=x.get("rooms"))
        ws.cell(row=r, column=4, value=(x.get("geo") or x.get("street") or "")[:80])
        c = ws.cell(row=r, column=5, value="відкрити")
        c.hyperlink, c.font = x.get("url"), LINK_FONT

    if drops:
        r += 2
        ws.cell(row=r, column=1, value="Подешевшали").font = SUB_FONT
        r += 1
        _header(ws, r, [f"Було, {sym}", f"Стало, {sym}", "Знижка", "Адреса", "Посилання"])
        for x in sorted(drops, key=lambda x: x[money] / x["first_" + money])[:15]:
            r += 1
            ws.cell(row=r, column=1, value=x["first_" + money]).number_format = "#,##0"
            ws.cell(row=r, column=2, value=x[money]).number_format = "#,##0"
            ws.cell(row=r, column=3, value=x[money] / x["first_" + money] - 1).number_format = "0.0%"
            ws.cell(row=r, column=4, value=(x.get("geo") or "")[:80])
            c = ws.cell(row=r, column=5, value="відкрити")
            c.hyperlink, c.font = x.get("url"), LINK_FONT
    return max(r + 3, r0 + 30)


def build_report(archive, searches: list, out_path: str) -> dict:
    """searches — список пошуків (dict з id, name, deal). Повертає коротку статистику."""
    wb = Workbook()
    an = wb.active
    an.title = "Аналітика"
    an.column_dimensions["A"].width = 32
    for col, w in zip("BCDE", (14, 14, 40, 14)):
        an.column_dimensions[col].width = w
    used, row, stats = {"Аналітика"}, 1, {}
    for s in searches:
        sid = s["id"]
        rows = archive.rows(sid)
        if not rows:
            continue
        title = _sheet_title(s["name"], used)
        listing_sheet(wb.create_sheet(title), s["name"], rows, s.get("deal", "sale"))
        row = analytics_block(an, row, s["name"], rows, s.get("deal", "sale"), title)
        stats[s["name"]] = len(rows)
    if not stats:
        an["A1"] = "Архів поки порожній — оголошення з'являться після перших перевірок."
    for ws in wb.worksheets:      # друк: альбомна орієнтація, усе по ширині сторінки
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.save(out_path)
    return stats
