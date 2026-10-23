"""Excel-звіт по знайдених авто: аркуш на пошук з посиланнями + «Аналітика»."""
import re
import statistics
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .caption import DRIVE_UA, FUEL_UA, GEAR_UA

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE = Font(bold=True, size=14)
SUB = Font(bold=True, size=11, color="1F4E78")
LINK = Font(color="0563C1", underline="single")
GOOD = Font(bold=True, color="2E7D32")
DROP_FILL = PatternFill("solid", fgColor="E2EFDA")

COLS = ["Знайдено", "Ціна $", "vs ринок", "Марка", "Модель", "Рік", "Пробіг, тис", "Двигун", "Пальне", "КПП", "Привід",
        "Місто", "Розмитн.", "ДТП", "Зміна ціни", "Джерело", "Посилання", "VIN", "Опис"]
WIDTH = [16, 10, 9, 14, 16, 7, 11, 8, 12, 10, 10, 16, 9, 7, 11, 18, 11, 20, 60]


def _med(v):
    v = [x for x in v if x is not None]
    return round(statistics.median(v)) if v else None


def _head(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _title(name, used):
    t = re.sub(r"[\[\]\*\?/\\:]", " ", name)[:28].strip() or "Пошук"
    base, n = t, 2
    while t in used:
        t = f"{base[:25]} {n}"
        n += 1
    used.add(t)
    return t


def list_sheet(ws, name, rows):
    ws.append([name])
    ws["A1"].font = TITLE
    ws.append([f"Авто в архіві: {len(rows)} · сформовано {datetime.now():%d.%m.%Y %H:%M}"])
    _head(ws, 3, COLS)
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in rows:
        delta = (r["price_usd"] - r["first_price_usd"]) if r.get("price_usd") and r.get("first_price_usd") else 0
        ws.append([
            r.get("first_seen"), r.get("price_usd"),
            (r["market_pct"] / 100) if r.get("market_pct") is not None else None,
            r.get("brand"), r.get("model"), r.get("year"), r.get("mileage"),
            float(r["engine"]) if r.get("engine") else None,
            FUEL_UA.get(r.get("fuel") or "", r.get("fuel") or ""), GEAR_UA.get(r.get("gearbox") or "", ""),
            DRIVE_UA.get(r.get("drive") or "", ""), r.get("city"),
            {"1": "так", "0": "ні"}.get(str(r.get("customs_ok")), ""), {"1": "так", "0": "ні"}.get(str(r.get("damaged")), ""),
            f"{delta:+,}$".replace(",", " ") if delta else "", r.get("source"), "відкрити", r.get("vin"),
            " ".join((r.get("text") or "").split())[:300],
        ])
        i = ws.max_row
        c = ws.cell(row=i, column=COLS.index("Посилання") + 1)
        c.hyperlink, c.font = r.get("url"), LINK
        pc = ws.cell(row=i, column=3)
        pc.number_format = "+0%;-0%;0%"
        if isinstance(pc.value, (int, float)) and pc.value <= -0.1:
            pc.font = GOOD
        ws.cell(row=i, column=2).number_format = "#,##0"
        if delta < 0:
            for k in range(1, len(COLS) + 1):
                ws.cell(row=i, column=k).fill = DROP_FILL
    ws.freeze_panes = "B4"
    if rows:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(COLS))}{ws.max_row}"


def analytics(ws, r0, name, rows, link_sheet):
    ws.cell(row=r0, column=1, value=name).font = TITLE
    c = ws.cell(row=r0, column=6, value="→ усі авто")
    c.hyperlink, c.font = f"#'{link_sheet}'!A1", LINK
    prices = [x["price_usd"] for x in rows if x.get("price_usd")]
    r = r0 + 1
    for k, v in [("Авто", len(rows)), ("Медіанна ціна, $", _med(prices)), ("Мінімальна, $", min(prices) if prices else None),
                 ("Максимальна, $", max(prices) if prices else None),
                 ("Медіанний пробіг, тис км", _med([x.get("mileage") for x in rows])),
                 ("Дешевше ринку на 10%+", sum(1 for x in rows if (x.get("market_pct") or 0) <= -10))]:
        ws.cell(row=r, column=1, value=k)
        cell = ws.cell(row=r, column=2, value=v)
        cell.number_format, cell.font = "#,##0", Font(bold=True)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="За моделлю і роком").font = SUB
    r += 1
    _head(ws, r, ["Модель, рік", "Авто", "Медіана, $", "Мін., $", "Пробіг, тис"])
    g = defaultdict(list)
    for x in rows:
        g[f"{x.get('brand') or ''} {x.get('model') or ''} {x.get('year') or ''}".strip()].append(x)
    first = r + 1
    for k, v in sorted(g.items(), key=lambda kv: (-len(kv[1]), kv[0]))[:25]:
        r += 1
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=len(v))
        ws.cell(row=r, column=3, value=_med([x.get("price_usd") for x in v])).number_format = "#,##0"
        ws.cell(row=r, column=4, value=min((x["price_usd"] for x in v if x.get("price_usd")), default=None)).number_format = "#,##0"
        ws.cell(row=r, column=5, value=_med([x.get("mileage") for x in v]))
    last = r

    r += 2
    ws.cell(row=r, column=1, value="За роком випуску").font = SUB
    r += 1
    _head(ws, r, ["Рік", "Авто", "Медіана, $", "Пробіг, тис"])
    by_year = defaultdict(list)
    for x in rows:
        if x.get("year"):
            by_year[x["year"]].append(x)
    y_first = r + 1
    for y in sorted(by_year):
        r += 1
        v = by_year[y]
        ws.cell(row=r, column=1, value=y)
        ws.cell(row=r, column=2, value=len(v))
        ws.cell(row=r, column=3, value=_med([x.get("price_usd") for x in v])).number_format = "#,##0"
        ws.cell(row=r, column=4, value=_med([x.get("mileage") for x in v]))
    y_last = r
    if y_last >= y_first:
        ch = BarChart()
        ch.title, ch.height, ch.width = "Медіанна ціна за роком, $", 7, 14
        ch.add_data(Reference(ws, min_col=3, min_row=y_first - 1, max_row=y_last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=y_first, max_row=y_last))
        ch.legend = None
        ws.add_chart(ch, f"H{r0 + 1}")

    r += 2
    ws.cell(row=r, column=1, value="Топ-10 найвигідніших відносно ринку").font = SUB
    r += 1
    _head(ws, r, ["Авто", "Ціна, $", "vs ринок", "Пробіг", "Посилання"])
    best = sorted([x for x in rows if x.get("market_pct") is not None], key=lambda x: x["market_pct"])[:10]
    for x in best:
        r += 1
        ws.cell(row=r, column=1, value=f"{x.get('brand')} {x.get('model')} {x.get('year')}")
        ws.cell(row=r, column=2, value=x.get("price_usd")).number_format = "#,##0"
        pc = ws.cell(row=r, column=3, value=x["market_pct"] / 100)
        pc.number_format = "+0%;-0%;0%"
        ws.cell(row=r, column=4, value=x.get("mileage"))
        c = ws.cell(row=r, column=5, value="відкрити")
        c.hyperlink, c.font = x.get("url"), LINK

    r += 2
    ws.cell(row=r, column=1, value="За джерелом").font = SUB
    r += 1
    _head(ws, r, ["Джерело", "Авто", "Медіана, $"])
    src = defaultdict(list)
    for x in rows:
        src[x.get("source") or "?"].append(x)
    for k, v in sorted(src.items(), key=lambda kv: -len(kv[1])):
        r += 1
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=len(v))
        ws.cell(row=r, column=3, value=_med([x.get("price_usd") for x in v])).number_format = "#,##0"
    return max(r + 3, r0 + 30)


def build_report(archive, searches, out_path):
    wb = Workbook()
    an = wb.active
    an.title = "Аналітика"
    for col, w in zip("ABCDE", (34, 10, 12, 10, 12)):
        an.column_dimensions[col].width = w
    used, row, stats = {"Аналітика"}, 1, {}
    for s in searches:
        rows = archive.rows(s["id"])
        if not rows:
            continue
        t = _title(s["name"], used)
        list_sheet(wb.create_sheet(t), s["name"], rows)
        row = analytics(an, row, s["name"], rows, t)
        stats[s["name"]] = len(rows)
    if not stats:
        an["A1"] = "Архів поки порожній."
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.save(out_path)
    return stats
